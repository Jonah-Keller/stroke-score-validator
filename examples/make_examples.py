"""Generate SYNTHETIC example sheets for validator.html, plus the expected answers.

Every value here is fake: drawn from a seeded random generator, with no link to
any real patient. Each sheet renders the same kind of synthetic cohort in a
different "site style" (headers, codings, date formats, file type) so the
column/value detection can be tested end to end (tests/test_examples.js).

For each sheet, examples/expected/<name>.json holds
  mapping       canonical variable -> the header the tool should end up using
  user_actions  what a real person must do in the column matcher (prompt answers,
                per-value codes the tool cannot know, e.g. REDCap numeric codes)
  prompts       which required questions should appear
  checks        known aggregate answers, computed here from the known truth
                (independently of the JavaScript engine)

Run:  /Users/jonahkeller/NeuroVERT/causal_ml/.venv/bin/python examples/make_examples.py
"""
from __future__ import annotations

import json
import math
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
EXP = HERE / "expected"
N = 80
TICI_GRADE = {0: "0", 1: "1", 2: "2a", 3: "2b", 4: "2c", 5: "3"}       # 0-5 code -> mTICI grade
RACES = ["White", "Black or African American", "Asian", "Other", "Unknown/declined"]
OCC = ["ICA", "M1", "M2", "Distal/MeVO (M3/M4/ACA/PCA)", "Posterior (basilar/vertebral)"]
EXCEL_EPOCH = datetime(1899, 12, 30)


# ---------------------------------------------------------------- truth cohort
def make_truth(seed: int) -> list[dict]:
    rng = np.random.default_rng(seed)
    start = datetime(2021, 3, 1)
    span_days = (datetime(2024, 8, 31) - start).days
    rows = []
    for i in range(N):
        age = int(np.clip(rng.normal(68, 13), 22, 97))
        nihss = int(np.clip(rng.normal(15, 6), 1, 36))
        tici = int(rng.choice([0, 1, 2, 3, 4, 5], p=[.06, .04, .12, .38, .15, .25]))
        passes = int(rng.choice([1, 2, 3, 4, 5], p=[.45, .25, .15, .1, .05]))
        premrs = int(rng.choice([0, 1, 2, 3, 4], p=[.55, .2, .12, .1, .03]))
        amb = int(rng.random() < (0.08 if premrs < 3 else 0.7))
        cancer = int(rng.random() < .13)
        dm = int(rng.random() < .3)
        smoke = int(rng.choice([0, 1, 2], p=[.5, .25, .25]))
        # outcome loosely tied to severity
        lin = 0.09 * nihss + 0.04 * (age - 68) - 0.35 * tici + 1.0 * cancer + 0.5 * premrs
        mrs90 = int(np.clip(round(2.4 + lin + rng.normal(0, 1.4)), 0, 6))
        died = int(mrs90 == 6 and rng.random() < .5)
        punct = start + timedelta(days=int(rng.integers(0, span_days)), hours=int(rng.integers(0, 24)),
                                  minutes=int(rng.integers(0, 60)))
        l2p = int(np.clip(rng.lognormal(math.log(330), 0.6), 45, 2800))
        rows.append(dict(
            age=age, sex="Female" if rng.random() < .52 else "Male",
            race=str(rng.choice(RACES, p=[.62, .25, .05, .04, .04])),
            eth=str(rng.choice(["Hispanic or Latino", "Not Hispanic or Latino", "Unknown"], p=[.15, .8, .05])),
            nihss=nihss, premrs=premrs, amb=amb, cancer=cancer, dm=dm, smoke=smoke, tici=tici, passes=passes,
            mrs90=mrs90, died=died, ivt=int(rng.random() < .35), prior=int(rng.random() < .16),
            occ=str(rng.choice(OCC, p=[.2, .42, .2, .1, .08])),
            punct=punct, lkn=punct - timedelta(minutes=l2p), l2p=l2p,
        ))
    return rows


# --------------------------------------------------------- expected aggregates
def q7(vals):
    v = sorted(vals)
    if not v:
        return None
    h = (len(v) - 1) * 0.5
    lo, hi = math.floor(h), math.ceil(h)
    return round(v[lo] + (h - lo) * (v[hi] - v[lo]), 2)


def checks_from(coded: list[dict], has: set[str]) -> dict:
    """Known answers from the coded truth (what the tool should produce)."""
    def cnt(f):
        return sum(1 for r in coded if f(r))
    out = {"n_patients": len(coded)}
    if "nihss" in has:
        out["nihss_median"] = q7([r["nihss"] for r in coded if r.get("nihss") is not None])
    if "age" in has:
        out["age_median"] = q7([r["age"] for r in coded if r.get("age") is not None])
        out["age_n"] = cnt(lambda r: r.get("age") is not None)
    if "mrs90" in has:
        def m(r):
            v = r.get("mrs90")
            if v is None and "died" in has and r.get("died") == 1:
                return 6
            return v
        out["mrs90_n"] = cnt(lambda r: m(r) is not None)
        out["mrs90_0_2"] = cnt(lambda r: m(r) is not None and m(r) <= 2)
        out["mrs90_6"] = cnt(lambda r: m(r) == 6)
    if "tici" in has:
        out["tici_n"] = cnt(lambda r: r.get("tici") is not None)
        out["tici_ge_2b"] = cnt(lambda r: r.get("tici") is not None and r["tici"] >= 3)
    if "sex" in has:
        out["female"] = cnt(lambda r: r.get("sex") == "Female")
    if "race" in has:
        out["race_white"] = cnt(lambda r: r.get("race") == "White")
        out["race_black"] = cnt(lambda r: r.get("race") == "Black or African American")
    if "eth" in has:
        out["hispanic"] = cnt(lambda r: r.get("eth") == "Hispanic or Latino")
    if "ivt" in has:
        out["ivt_yes"] = cnt(lambda r: r.get("ivt") == 1)
    if "occ" in has:
        out["occ_m1"] = cnt(lambda r: r.get("occ") == "M1")
    if "amb" in has:
        out["nonambulatory"] = cnt(lambda r: r.get("amb") == 1)
    if "smoke" in has:
        out["smoke_never"] = cnt(lambda r: r.get("smoke") == 0)
        out["smoke_current"] = cnt(lambda r: r.get("smoke") == 1)
        out["smoke_former"] = cnt(lambda r: r.get("smoke") == 2)
    if "l2p" in has:
        ok = [r["l2p"] for r in coded if r.get("l2p") is not None and 0 <= r["l2p"] <= 1440 * 7]
        out["l2p_median"] = q7(ok)
        out["l2p_implausible"] = cnt(lambda r: r.get("l2p") is not None and not (0 <= r["l2p"] <= 1440 * 7))
    if "proc" in has:
        ds = [r["proc"] for r in coded if r.get("proc") is not None]
        out["first_month"] = min(ds).strftime("%Y-%m")
        out["last_month"] = max(ds).strftime("%Y-%m")
    peri = ["nihss", "age", "cancer", "passes", "tici"]
    pre = ["nihss", "cancer", "dm", "amb", "premrs", "smoke"]
    if set(peri) <= has:
        out["peri_n_scored"] = cnt(lambda r: all(r.get(k) is not None for k in peri))
    if set(pre) <= has:
        sc = [2 * r["nihss"] + 22 * r["cancer"] + 7 * r["dm"] + 7 * r["amb"] + 6 * r["premrs"] + 2 * r["smoke"]
              for r in coded if all(r.get(k) is not None for k in pre)]
        out["preop_n_scored"] = len(sc)
        out["preop_median"] = q7(sc)
        out["preop_ge_42"] = sum(1 for s in sc if s >= 42)
    return out


def write_expected(name, mapping, actions, prompts, coded, has, tici_scheme, extra=None):
    EXP.mkdir(exist_ok=True)
    d = {"file": name, "synthetic": True, "mapping": mapping, "user_actions": actions,
         "prompts": prompts, "tici_scheme": tici_scheme, "checks": checks_from(coded, has)}
    if extra:
        d.update(extra)
    (EXP / (Path(name).stem + ".json")).write_text(json.dumps(d, indent=2) + "\n")


def blank_some(rng, vals, frac, token=""):
    return [token if rng.random() < frac else v for v in vals]


# ------------------------------------------------------------------ 1. study style
def study_style():
    T = make_truth(101)
    df = pd.DataFrame({
        "age": [r["age"] for r in T], "nihss": [r["nihss"] for r in T], "tici": [r["tici"] for r in T],
        "number_of_passes": [r["passes"] for r in T], "hx_cancer": [r["cancer"] for r in T],
        "hx_dm2": [r["dm"] for r in T], "pre_stroke_ambulatory_status": [r["amb"] for r in T],
        "mrs_pre_stroke": [r["premrs"] for r in T], "hx_smoking": [r["smoke"] for r in T],
        "mrs_90day": [r["mrs90"] for r in T], "mortality_hospital": [r["died"] for r in T],
        "sex": [r["sex"] for r in T], "race": [r["race"] for r in T], "ethnicity": [r["eth"] for r in T],
        "iv_thrombolysis": [r["ivt"] for r in T], "prior_stroke": [r["prior"] for r in T],
        "occlusion_site": [r["occ"].split(" ")[0].replace("Distal/MeVO", "M3").replace("Posterior", "Basilar") for r in T],
        "lkn_datetime": [r["lkn"].strftime("%Y-%m-%d %H:%M") for r in T],
        "puncture_datetime": [r["punct"].strftime("%Y-%m-%d %H:%M") for r in T],
        "procedure_date": [r["punct"].strftime("%Y-%m-%d") for r in T],
    })
    df.to_csv(HERE / "study_style.csv", index=False)
    mapping = {c: c for c in df.columns}
    coded = [dict(r, proc=r["punct"]) for r in T]
    write_expected("study_style.csv", mapping,
                   {"choices": {"ambulatory_yes_means": "dependent"}},
                   ["pre_stroke_ambulatory_status"], coded,
                   {"age", "nihss", "tici", "passes", "cancer", "dm", "amb", "premrs", "smoke", "mrs90", "died",
                    "sex", "race", "eth", "ivt", "occ", "l2p", "proc"}, "study05")


# --------------------------------------------------------------- 2. registry style
REG_HEADERS = {"age": "Age", "nihss": "Admission NIHSS", "hx_cancer": "Hx Malignancy", "hx_dm2": "DM",
               "mrs_pre_stroke": "Pre-morbid mRS", "pre_stroke_ambulatory_status": "Ambulation prior to stroke",
               "hx_smoking": "Tobacco use", "tici": "mTICI", "number_of_passes": "# of passes",
               "mrs_90day": "90 day mRS", "sex": "Gender", "race": "Race", "ethnicity": "Ethnicity",
               "iv_thrombolysis": "tPA given", "lkn_datetime": "LKW", "puncture_datetime": "Groin puncture time",
               "procedure_date": "Procedure Date", "occlusion_site": "Occlusion location", "prior_stroke": "Prior CVA"}
OCC_TXT = {"ICA": "ICA", "M1": "M1", "M2": "M2", "Distal/MeVO (M3/M4/ACA/PCA)": "M3",
           "Posterior (basilar/vertebral)": "Basilar"}


def registry_frame(T):
    yn = lambda v: "Yes" if v else "No"
    return {
        "Age": [r["age"] for r in T], "Admission NIHSS": [r["nihss"] for r in T],
        "Hx Malignancy": [yn(r["cancer"]) for r in T], "DM": [yn(r["dm"]) for r in T],
        "Pre-morbid mRS": [r["premrs"] for r in T],
        "Ambulation prior to stroke": ["Needs assistance" if r["amb"] else "Independent" for r in T],
        "Tobacco use": [["Never", "Current", "Former"][r["smoke"]] for r in T],
        "mTICI": [TICI_GRADE[r["tici"]] for r in T], "# of passes": [r["passes"] for r in T],
        "90 day mRS": [r["mrs90"] for r in T], "Gender": [r["sex"] for r in T],
        "Race": [r["race"] if r["race"] != "Unknown/declined" else "Declined" for r in T],
        "Ethnicity": [r["eth"] for r in T], "tPA given": [yn(r["ivt"]) for r in T],
        "LKW": [r["lkn"] for r in T], "Groin puncture time": [r["punct"] for r in T],
        "Procedure Date": [r["punct"].replace(hour=0, minute=0) for r in T],
        "Occlusion location": [OCC_TXT[r["occ"]] for r in T], "Prior CVA": [yn(r["prior"]) for r in T],
    }


def registry_style():
    T = make_truth(202)
    cols = registry_frame(T)
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Registry"
    hdr = list(cols)
    ws.append(hdr)
    for i in range(N):
        ws.append([cols[h][i] for h in hdr])
    for c in ("O", "P"):
        for cell in ws[c][1:]:
            cell.number_format = "m/d/yyyy h:mm"
    for cell in ws["Q"][1:]:
        cell.number_format = "m/d/yyyy"
    wb.save(HERE / "registry_style.xlsx")
    coded = [dict(r, proc=r["punct"]) for r in T]
    write_expected("registry_style.xlsx", dict(REG_HEADERS), {"choices": {}}, [], coded,
                   {"age", "nihss", "tici", "passes", "cancer", "dm", "amb", "premrs", "smoke", "mrs90",
                    "sex", "race", "eth", "ivt", "occ", "l2p", "proc"}, "mtici")
    # 7. the same sheet as legacy BIFF .xls (the page cannot read BIFF; it must say so clearly)
    import xlwt
    wb2 = xlwt.Workbook(); s2 = wb2.add_sheet("Registry")
    dfmt = xlwt.easyxf(num_format_str="M/D/YYYY h:mm")
    for j, h in enumerate(hdr):
        s2.write(0, j, h)
        for i in range(N):
            v = cols[h][i]
            if isinstance(v, datetime):
                s2.write(i + 1, j, v, dfmt)
            else:
                s2.write(i + 1, j, v)
    wb2.save(str(HERE / "registry_style.xls"))
    (EXP / "registry_style_xls.json").write_text(json.dumps(
        {"file": "registry_style.xls", "synthetic": True,
         "expect_error": "old-format Excel (.xls) is not readable in the browser: re-save as .xlsx or .csv"}, indent=2) + "\n")


# --------------------------------------------------------------- 3. messy export
def messy_export():
    T = make_truth(303)
    rng = np.random.default_rng(3030)
    race_code = {"White": "5", "Black or African American": "3", "Asian": "2", "Other": "6", "Unknown/declined": "7"}
    occ_txt = {"ICA": ["L ICA", "R ICA ", "T-ICA"], "M1": ["L M1", "R M1", "m1 "], "M2": ["L M2", "R M2"],
               "Distal/MeVO (M3/M4/ACA/PCA)": ["M3", "ACA", "PCA"], "Posterior (basilar/vertebral)": ["BASILAR", "Vertebral"]}
    coded, raw = [], []
    for i, r in enumerate(T):
        c = dict(r, proc=r["punct"])          # no procedure-date column: window comes from puncture time
        age = str(r["age"]) if rng.random() > .05 else ""
        if age == "":
            c["age"] = None
        mr = r["mrs90"]
        u = rng.random()
        if u < .06:
            mrs_txt, c["mrs90"] = "", None
        elif u < .1:
            mrs_txt, c["mrs90"] = "UNK", None
        elif mr == 6 and rng.random() < .5:
            mrs_txt = "dead"
        elif rng.random() < .25:
            mrs_txt = "mRS %d" % mr
        else:
            mrs_txt = str(mr)
        canc = ("1" if r["cancer"] else "0")
        if rng.random() < .05:
            canc, c["cancer"] = "unknown", None
        tic = TICI_GRADE[r["tici"]].upper() if rng.random() < .3 else TICI_GRADE[r["tici"]]
        if rng.random() < .04:
            tic, c["tici"] = "N/A", None
        # smoking as yes/no only: yes = ever smoked -> coded 1 (current) once the user says "ever"
        smk = "yes" if r["smoke"] in (1, 2) else "no"
        c["smoke"] = 1 if r["smoke"] in (1, 2) else 0
        l2p = r["l2p"]
        lkn = r["lkn"]
        if i in (5, 41):                      # data-entry errors: LKN after puncture
            lkn = r["punct"] + timedelta(minutes=30)
            l2p = -30
        c["l2p"] = l2p
        raw.append({
            "Record ID": "SYN-%04d" % (i + 1),
            "Age (yrs)": age,
            "Sex (1=M, 2=F)": "1" if r["sex"] == "Male" else "2",
            "race": race_code[r["race"]],
            "NIHSS_admit": str(r["nihss"]),
            "premorbid_mrs": str(r["premrs"]),
            "ambulatory": "0" if r["amb"] else "1",            # REVERSED: 1 = ambulatory
            "cancer_hx": canc,
            "diabetes": "Y" if r["dm"] else "N",
            "smoker": smk + (" " if rng.random() < .2 else ""),
            "tici_final": tic,
            "num_passes": str(r["passes"]),
            "mrs_90d": mrs_txt,
            "ivt": "1" if r["ivt"] else "0",
            "occlusion_vessel": str(rng.choice(occ_txt[r["occ"]])),
            "lkw_dttm": lkn.strftime("%-m/%-d/%Y %H:%M"),
            "groin_dttm": r["punct"].strftime("%-m/%-d/%Y %H:%M"),
        })
        coded.append(c)
    pd.DataFrame(raw).to_csv(HERE / "messy_export.csv", index=False)
    mapping = {"age": "Age (yrs)", "sex": "Sex (1=M, 2=F)", "race": "race", "nihss": "NIHSS_admit",
               "mrs_pre_stroke": "premorbid_mrs", "pre_stroke_ambulatory_status": "ambulatory",
               "hx_cancer": "cancer_hx", "hx_dm2": "diabetes", "hx_smoking": "smoker", "tici": "tici_final",
               "number_of_passes": "num_passes", "mrs_90day": "mrs_90d", "iv_thrombolysis": "ivt",
               "occlusion_site": "occlusion_vessel", "lkn_datetime": "lkw_dttm", "puncture_datetime": "groin_dttm"}
    actions = {"choices": {"ambulatory_yes_means": "independent", "smoking_yes_means": "ever"},
               "value_codes": {"sex": {"1": "Male", "2": "Female"},
                               "race": {"5": "White", "3": "Black or African American", "2": "Asian",
                                        "6": "Other", "7": "Unknown/declined"}}}
    write_expected("messy_export.csv", mapping, actions,
                   ["pre_stroke_ambulatory_status", "hx_smoking"], coded,
                   {"age", "nihss", "tici", "passes", "cancer", "dm", "amb", "premrs", "smoke", "mrs90",
                    "sex", "race", "ivt", "occ", "l2p", "proc"}, "mtici",
                   {"enrollment_source": "puncture_datetime"})


# ------------------------------------------------- 4. European semicolon / D/M/YYYY
def european():
    T = make_truth(404)
    raw = []
    coded = []
    for r in T:
        raw.append({
            "Alter": "", "age_years": r["age"], "NIHSS": r["nihss"], "TICI": TICI_GRADE[r["tici"]],
            "passes": r["passes"], "cancer": "ja" if False else ("1" if r["cancer"] else "0"),
            "diabetes": "1" if r["dm"] else "0", "premorbid mRS": r["premrs"],
            "smoking": ["never", "current", "former"][r["smoke"]],
            "mRS 90 days": r["mrs90"], "sex": "F" if r["sex"] == "Female" else "M",
            "thrombectomy date": r["punct"].strftime("%d/%m/%Y"),
            "LKN to puncture (min)": ("%.1f" % (r["l2p"] + 0.5)).replace(".", ","),
        })
        coded.append(dict(r, l2p=r["l2p"] + 0.5, proc=r["punct"], amb=None))
    df = pd.DataFrame(raw).drop(columns=["Alter"])
    df.to_csv(HERE / "semicolon_european.csv", sep=";", index=False)
    mapping = {"age": "age_years", "nihss": "NIHSS", "tici": "TICI", "number_of_passes": "passes",
               "hx_cancer": "cancer", "hx_dm2": "diabetes", "mrs_pre_stroke": "premorbid mRS",
               "hx_smoking": "smoking", "mrs_90day": "mRS 90 days", "sex": "sex",
               "procedure_date": "thrombectomy date", "lkn_to_puncture_min": "LKN to puncture (min)"}
    write_expected("semicolon_european.csv", mapping, {"choices": {}}, [], coded,
                   {"age", "nihss", "tici", "passes", "cancer", "smoke", "mrs90", "sex", "l2p", "proc"}, "mtici",
                   {"date_order": {"procedure_date": "DM"}})


# ------------------------------------------- 5. Excel serials + decoy columns + sheets
def excel_serials():
    T = make_truth(505)
    rng = np.random.default_rng(5050)
    from openpyxl import Workbook
    wb = Workbook()
    ws0 = wb.active; ws0.title = "Instructions"
    ws0.append(["This workbook is synthetic test data."]); ws0.append(["Data are on the next sheet."])
    ws = wb.create_sheet("Data")
    hdr = ["MRN", "Location", "Time", "Notes", "Patient age", "NIH Stroke Scale", "eTICI", "Passes",
           "Active cancer", "Diabetes mellitus", "Baseline mRS", "Smoking status", "mRS at 90 days",
           "Sex", "Last known well", "Arterial puncture", "Occluded vessel", "IV alteplase", "Discharge status"]
    ws.append(hdr)
    serial = lambda d: round((d - EXCEL_EPOCH).total_seconds() / 86400, 6)
    coded = []
    for i, r in enumerate(T):
        lkn = serial(r["lkn"]) if i % 3 else r["lkn"].strftime("%Y-%m-%d %H:%M")   # numbers and text mixed
        pun = serial(r["punct"]) if i % 4 else r["punct"].strftime("%Y-%m-%dT%H:%M:%S")
        ws.append(["SYN%07d" % (4000000 + i), str(rng.choice(["Main campus", "North tower", "ED"])),
                   "%02d:%02d" % (rng.integers(0, 24), rng.integers(0, 60)), "synthetic row",
                   r["age"], r["nihss"], TICI_GRADE[r["tici"]], r["passes"],
                   "TRUE" if r["cancer"] else "FALSE", "TRUE" if r["dm"] else "FALSE", r["premrs"],
                   ["Never smoker", "Current every day smoker", "Former smoker"][r["smoke"]], r["mrs90"],
                   r["sex"][0], lkn, pun, OCC_TXT[r["occ"]], "Y" if r["ivt"] else "N",
                   "Expired" if r["died"] else "Discharged"])
        coded.append(dict(r))
    wb.save(HERE / "excel_serials.xlsx")
    mapping = {"age": "Patient age", "nihss": "NIH Stroke Scale", "tici": "eTICI", "number_of_passes": "Passes",
               "hx_cancer": "Active cancer", "hx_dm2": "Diabetes mellitus", "mrs_pre_stroke": "Baseline mRS",
               "hx_smoking": "Smoking status", "mrs_90day": "mRS at 90 days", "sex": "Sex",
               "lkn_datetime": "Last known well", "puncture_datetime": "Arterial puncture",
               "occlusion_site": "Occluded vessel", "iv_thrombolysis": "IV alteplase",
               "mortality_hospital": "Discharge status"}
    write_expected("excel_serials.xlsx", mapping, {"choices": {}, "sheet": "Data"}, [], coded,
                   {"age", "nihss", "tici", "passes", "cancer", "smoke", "mrs90", "died", "sex", "occ", "ivt", "l2p"},
                   "mtici", {"must_not_map": ["MRN", "Location", "Time", "Notes"], "sheets": ["Instructions", "Data"]})


# ------------------------------------------------------------------- 6. minimal
def minimal():
    T = make_truth(606)
    df = pd.DataFrame({"NIHSS": [r["nihss"] for r in T], "cancer": [r["cancer"] for r in T],
                       "mRS90": [r["mrs90"] for r in T]})
    df.to_csv(HERE / "minimal.csv", index=False)
    write_expected("minimal.csv", {"nihss": "NIHSS", "hx_cancer": "cancer", "mrs_90day": "mRS90"},
                   {"choices": {}}, [], [dict(r) for r in T], {"nihss", "cancer", "mrs90"}, None)


if __name__ == "__main__":
    EXP.mkdir(exist_ok=True)
    for f in (study_style, registry_style, messy_export, european, excel_serials, minimal):
        f()
    print("wrote", sorted(p.name for p in HERE.iterdir() if p.is_file() and p.suffix != ".py"))
